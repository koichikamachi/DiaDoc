"""診断経緯の文章化（2026-09-30）：構造化レポートを材料に、二部構成の報告書の下書きを書く。"""

from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest

from core import narrative
from core.graph import DebateSession
from core.runs import latest_run

ROOT = Path(__file__).resolve().parents[1]
CRISIS = "C002_sample_crisis"


@pytest.fixture()
def data(tmp_path, monkeypatch):
    shutil.copytree(ROOT / "data", tmp_path / "data")
    monkeypatch.setenv("DBD_DATA_DIR", str(tmp_path / "data"))
    return tmp_path / "data"


def _finished():
    s = DebateSession(latest_run(CRISIS))
    for _ in range(40):
        if s.finished:
            return latest_run(CRISIS)
        s.step()
    raise AssertionError


class _Writer:
    name = "fake"

    def __init__(self, text):
        self.text, self.seen = text, None

    def write(self, system, user):
        self.seen = (system, user)
        return self.text


def test_only_finished_runs_can_be_narrated(data):
    with pytest.raises(ValueError, match="診断完了"):
        narrative.generate(latest_run(CRISIS), writer=_Writer("x"))


def test_mock_draft_has_both_parts_and_all_headings(data):
    run = _finished()
    nar = narrative.generate(run)
    assert nar.engine == "mock"
    t = nar.text
    assert "# 第1部：第1次経営診断 審理経緯レポート" in t and "# 第2部：金融機関提出用 経営診断サマリー" in t
    for h in narrative.PART1_SECTIONS + narrative.PART2_SECTIONS:
        assert f"## {h}" in t
    assert nar.unverified_numbers == []                                  # 材料の行を並べただけなので、数字はすべて材料にある
    assert nar.document().startswith(narrative.DRAFT_NOTE)


def test_prompt_passes_the_report_as_data_and_forbids_inventing(data):
    run = _finished()
    w = _Writer("# 第1部\n本文")
    narrative.generate(run, writer=w)
    system, user = w.seen
    assert "記録にない数字・銀行名・人名・日付・計画値を作らない" in system and "（記録なし）" in system
    assert "## 2. 資金不足額の確定" in system and "## 金融機関へのお願い" in system
    assert "<untrusted_document" in user and "経営診断レポート：C002" in user
    assert "untrusted_document> の中身は" in system                        # 資料の中の命令に従わない


def test_numbers_not_in_the_report_are_flagged(data):
    run = _finished()
    nar = narrative.generate(run, writer=_Writer("必要CFは年122,090千円、改善効果は年99,999千円と見込む。△21,290。"))
    assert nar.unverified_numbers == ["99,999"]
    assert "要確認" in nar.document() and "99,999" in nar.document()


def test_saved_with_audit_trail_and_reused(data):
    run = _finished()
    nar = narrative.generate(run)
    again = narrative.load(run)
    assert again.text == nar.text and again.report_hash == nar.report_hash
    assert run.read_text(narrative.FILE_MD).startswith(narrative.DRAFT_NOTE)
    log = run.read("audit_log.json")[-1]
    assert log["action"] == "診断経緯の文章化（下書きを生成）" and log["actor"] == "システム（文章化）"


def test_gemini_writer_sends_system_instruction():
    from google.genai import types

    class _Resp:
        text = "# 第1部\n本文"

    class _Models:
        def generate_content(self, model, contents, config):
            assert isinstance(config, types.GenerateContentConfig) and config.system_instruction == "SYS"
            _Models.seen = (model, contents)
            return _Resp()

    class _Client:
        models = _Models()

    w = narrative.GeminiWriter(client=_Client(), model="m")
    assert w.write("SYS", "USER") == "# 第1部\n本文" and _Models.seen == ("m", "USER") and w.name == "gemini:m"


def test_docx_uses_meiryo_headings_table_header_and_page_number():
    from docx import Document

    from core.docx_export import markdown_to_docx

    md = "# 第1部\n## 1. 現状\n本文の**太字**です。\n- 箇条\n| 施策 | 額 |\n|---|---:|\n| 労務費 | 1,500 |\n"
    doc = Document(io.BytesIO(markdown_to_docx(md, "第1次分析 診断経緯")))
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads == [("Heading 1", "第1部"), ("Heading 2", "1. 現状")]
    assert doc.styles["Normal"].font.name == "Meiryo" and doc.styles["Normal"].font.size.pt == 10.5
    assert doc.styles["Heading 1"].font.size.pt == 13 and doc.styles["Heading 2"].font.size.pt == 12
    assert any(r.bold for p in doc.paragraphs for r in p.runs if r.text == "太字")
    t = doc.tables[0]
    assert t.cell(1, 0).text == "労務費" and t.cell(1, 0).paragraphs[0].style.name == "標準表内文字"
    sec = doc.sections[0]
    assert sec.header.paragraphs[0].text == "第1次分析 診断経緯"
    assert "PAGE" in sec.footer.paragraphs[0]._p.xml


def test_sidebar_button_only_after_the_diagnosis_is_complete(data):
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "src/ui/app.py"), default_timeout=30)
    at.run()
    at.selectbox(key="company").set_value(CRISIS).run()
    assert at.button(key="narrative_btn").disabled
    _finished()
    at.run()
    btn = at.button(key="narrative_btn")
    assert not btn.disabled
    btn.click().run()
    assert not at.exception, at.exception
    assert narrative.load(latest_run(CRISIS)) is not None                  # 押すと下書きが作られて保存される


def test_prompt_forbids_writing_events_that_did_not_happen():
    """介入・確定が記録に「なし」なのに「介入を確認したところ」と書かせない（乙精機ケース3の下書きで実例）。"""
    from core.narrative import system_prompt

    p = system_prompt("1")
    assert "行われたかのような書き方をしない" in p and "介入と確定判断を行った）" not in p


def test_both_parts_open_with_a_summary():
    from core.narrative import PART1_SECTIONS, PART2_SECTIONS, SUMMARY, system_prompt

    assert PART1_SECTIONS[0] == PART2_SECTIONS[0] == SUMMARY == "要約"
    p = system_prompt("1")
    assert p.count("## 要約") == 2 and "論文の要約と同じ要領" in p
