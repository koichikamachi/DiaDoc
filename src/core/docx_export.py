"""Markdown（見出し・箇条書き・段落・太字・表）を Word（.docx）に書き出す。

書式（利用者の既定）：全文メイリオ。見出し1＝13pt、見出し2＝12pt、見出し3＝11pt太字。標準＝10.5pt・段落前0.5pt。
表の中は「標準表内文字」（メイリオ9pt・段落前後アキなし）。ヘッダーに標題（9pt）、フッターにページ番号（10pt）。
"""

from __future__ import annotations

import io
import re

FONT = "Meiryo"


def _font(obj, size: float | None = None, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    f = obj.font
    f.name = FONT
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rpr = obj.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)


def _styles(doc) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt, RGBColor

    normal = doc.styles["Normal"]
    _font(normal, 10.5)
    normal.paragraph_format.space_before = Pt(0.5)
    for name, size, bold in (("Heading 1", 13, True), ("Heading 2", 12, True), ("Heading 3", 11, True)):
        st = doc.styles[name]
        _font(st, size, bold)
        st.font.color.rgb = RGBColor(0, 0, 0)
    if "標準表内文字" not in [s.name for s in doc.styles]:
        t = doc.styles.add_style("標準表内文字", WD_STYLE_TYPE.PARAGRAPH)
        t.base_style = normal
        _font(t, 9)
        t.paragraph_format.space_before = Pt(0)
        t.paragraph_format.space_after = Pt(0)


def _runs(par, text: str, size: float | None = None) -> None:
    """**太字** だけを解釈して段落に文字を足す。"""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = par.add_run(part)
        _font(r, size, bold=True if i % 2 else None)


def _page_number(par) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = par.add_run()
    _font(run, 10)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def markdown_to_docx(md: str, title: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    _styles(doc)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    _runs(hp, title, 9)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_number(fp)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip() or ln.strip() == "---":
            i += 1
            continue
        if ln.lstrip().startswith("|"):   # 表：連続する | 行をまとめる
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncol)
                tbl.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c in range(ncol):
                        p = tbl.cell(r, c).paragraphs[0]
                        p.style = doc.styles["標準表内文字"]
                        _runs(p, cells[c] if c < len(cells) else "", 9)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            _runs(h, m.group(2))
        elif re.match(r"^\s*[-*]\s+", ln):
            depth = (len(ln) - len(ln.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if depth else "List Bullet")
            _runs(p, re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+[.．]\s+", ln):
            p = doc.add_paragraph(style="List Number")
            _runs(p, re.sub(r"^\s*\d+[.．]\s+", "", ln))
        elif ln.startswith(">"):
            p = doc.add_paragraph(style="Quote")
            _runs(p, ln.lstrip("> "))
        else:
            _runs(doc.add_paragraph(), ln)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
